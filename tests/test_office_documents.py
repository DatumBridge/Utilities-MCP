"""In-memory Word, Excel, and PowerPoint create/update."""

from __future__ import annotations

import base64
import io
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES
from pptx import Presentation

from app.core.exceptions import OfficeDocumentError
from app.services.office_documents import (
    MIME_DOCX,
    MIME_PPTX,
    MIME_XLSX,
    create_docx,
    create_pptx,
    create_xlsx,
    update_docx,
    update_pptx,
    update_xlsx,
)


def _docx_text(payload: str) -> str:
    document = Document(io.BytesIO(base64.b64decode(payload)))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def test_create_docx_title_paragraph_table_and_unicode():
    result = create_docx(
        file_name="Notes.docx",
        title="Tiếng Việt",
        paragraphs=[{"text": "Hello", "heading_level": 2}, {"text": "Body"}],
        tables=[{"headers": ["Item", "Qty"], "rows": [["Paper", "2"]]}],
    )
    assert result.file_name == "Notes.docx"
    assert result.mime_type == MIME_DOCX
    assert result.byte_length > 0
    assert result.replacements is None
    text = _docx_text(result.content_base64)
    assert "Tiếng Việt" in text
    assert "Hello" in text
    assert "Paper | 2" in text


def test_create_docx_requires_content():
    with pytest.raises(OfficeDocumentError) as caught:
        create_docx()
    assert caught.value.code == "VALIDATION_ERROR"
    assert caught.value.message == "Provide at least one paragraph, table, or title"


def test_update_docx_replaces_across_runs_then_appends():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Hel")
    paragraph.add_run("lo world")
    buffer = io.BytesIO()
    document.save(buffer)
    original = base64.b64encode(buffer.getvalue()).decode("ascii")
    result = update_docx(
        content_base64=original,
        file_name="edited.docx",
        replacements=[{"find": "Hello", "replace": "Xin chào"}],
        append_paragraphs=[{"text": "End"}],
    )
    assert result.replacements == 1
    text = _docx_text(result.content_base64)
    assert "Xin chào world" in text
    assert "End" in text


def test_update_docx_replace_miss_is_validation_error():
    created = create_docx(paragraphs=[{"text": "alpha"}])
    with pytest.raises(OfficeDocumentError) as caught:
        update_docx(content_base64=created.content_base64, replacements=[{"find": "missing", "replace": "x"}])
    assert caught.value.message == "replacement matched no text"


def test_update_docx_rejects_workbook():
    workbook = create_xlsx(sheets=[{"name": "Budget", "rows": [["1"]]}])
    with pytest.raises(OfficeDocumentError) as caught:
        update_docx(content_base64=workbook.content_base64, append_paragraphs=[{"text": "x"}])
    assert caught.value.code == "INVALID_OFFICE_FILE"
    assert caught.value.message == "Package is missing word/document.xml"


def test_create_xlsx_types_and_formula_text():
    result = create_xlsx(
        file_name="budget.xlsx",
        sheets=[
            {
                "name": "Budget",
                "headers": ["Item", "Qty", "Ok"],
                "rows": [["Paper", 2, True], ["=1+1", None, False]],
            }
        ],
    )
    assert result.mime_type == MIME_XLSX
    workbook = load_workbook(io.BytesIO(base64.b64decode(result.content_base64)))
    sheet = workbook["Budget"]
    assert sheet["A1"].value == "Item"
    assert sheet["B2"].value == 2
    assert sheet["C2"].value is True
    assert sheet["A3"].value == "=1+1"
    assert sheet["A3"].data_type == "s"
    assert sheet["B3"].value is None
    assert workbook.sheetnames == ["Budget"]


def test_update_xlsx_add_set_and_append():
    created = create_xlsx(sheets=[{"name": "Budget", "headers": ["Item"], "rows": [["Paper"]]}])
    result = update_xlsx(
        content_base64=created.content_base64,
        add_sheets=[{"name": "Notes", "rows": [["later"]]}],
        set_cells=[{"sheet": "Budget", "cell": "b2", "value": "Pen"}],
        append_rows=[{"sheet": "notes", "rows": [["more"]]}],
    )
    workbook = load_workbook(io.BytesIO(base64.b64decode(result.content_base64)))
    assert workbook["Budget"]["B2"].value == "Pen"
    assert workbook["Notes"]["A1"].value == "later"
    assert workbook["Notes"]["A2"].value == "more"


def test_update_xlsx_unknown_sheet_and_bad_cell():
    created = create_xlsx(sheets=[{"name": "Budget"}])
    with pytest.raises(OfficeDocumentError) as missing:
        update_xlsx(
            content_base64=created.content_base64,
            set_cells=[{"sheet": "Other", "cell": "A1", "value": 1}],
        )
    assert missing.value.message == "sheet not found: Other"
    with pytest.raises(OfficeDocumentError) as bad_cell:
        update_xlsx(
            content_base64=created.content_base64,
            set_cells=[{"sheet": "Budget", "cell": "A0", "value": 1}],
        )
    assert bad_cell.value.message == "cell must be an A1 reference such as B2"


def test_create_and_update_pptx():
    created = create_pptx(
        file_name="deck.pptx",
        slides=[{"title": "Kickoff", "bullets": ["One", "Two"], "notes": "Say hello"}],
    )
    assert created.mime_type == MIME_PPTX
    presentation = Presentation(io.BytesIO(base64.b64decode(created.content_base64)))
    assert len(presentation.slides) == 1
    texts = []
    for shape in presentation.slides[0].shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
    assert "Kickoff" in texts
    assert "One" in "\n".join(texts)
    assert presentation.slides[0].notes_slide.notes_text_frame.text == "Say hello"

    updated = update_pptx(
        content_base64=created.content_base64,
        replacements=[{"find": "One", "replace": "Một"}],
        set_slide_titles=[{"index": 0, "title": "Start"}],
        append_slides=[{"title": "Close", "bullets": ["Done"]}],
    )
    assert updated.replacements == 1
    presentation = Presentation(io.BytesIO(base64.b64decode(updated.content_base64)))
    assert len(presentation.slides) == 2
    title_shape = next(shape for shape in presentation.slides[0].shapes if shape.name == "Title")
    assert title_shape.text_frame.text == "Start"
    body = next(shape for shape in presentation.slides[0].shapes if shape.name == "Body")
    assert "Một" in body.text_frame.text


def test_rejects_bad_name_base64_and_empty_edit():
    with pytest.raises(OfficeDocumentError) as name:
        create_docx(title="T", file_name="../escape.docx")
    assert "file_name must be a base name ending with .docx" in name.value.message
    created = create_docx(title="T")
    with pytest.raises(OfficeDocumentError) as bad:
        update_docx(content_base64="!!!!", append_paragraphs=[{"text": "x"}])
    assert bad.value.code == "INVALID_BASE64"
    assert bad.value.message.startswith("Invalid Base64:")
    with pytest.raises(OfficeDocumentError) as wrong_package:
        update_xlsx(
            content_base64=created.content_base64,
            set_cells=[{"sheet": "Budget", "cell": "A1", "value": 1}],
        )
    assert wrong_package.value.code == "INVALID_OFFICE_FILE"
    assert wrong_package.value.message == "Package is missing xl/workbook.xml"
    with pytest.raises(OfficeDocumentError) as no_edit:
        update_pptx(content_base64=create_pptx(slides=[{"title": "A"}]).content_base64)
    assert no_edit.value.message == "Provide at least one of append_slides, replacements, or set_slide_titles"


def test_create_xlsx_keeps_sheet_and_sheet1_names():
    result = create_xlsx(
        sheets=[
            {"name": "Sheet", "rows": [["a"]]},
            {"name": "Sheet1", "rows": [["b"]]},
        ]
    )
    workbook = load_workbook(io.BytesIO(base64.b64decode(result.content_base64)))
    assert workbook.sheetnames == ["Sheet", "Sheet1"]
    assert workbook["Sheet"]["A1"].value == "a"
    assert workbook["Sheet1"]["A1"].value == "b"
    lower = create_xlsx(sheets=[{"name": "sheet", "rows": [["c"]]}])
    named = load_workbook(io.BytesIO(base64.b64decode(lower.content_base64)))
    assert named.sheetnames == ["sheet"]


def test_xlsx_save_does_not_create_temp_files(monkeypatch):
    import openpyxl.worksheet._writer as writer

    def forbid_temp(*_args, **_kwargs):
        raise AssertionError("worksheet XML must stay in memory")

    monkeypatch.setattr(writer, "create_temporary_file", forbid_temp)
    before = len(ALL_TEMP_FILES)
    create_xlsx(sheets=[{"name": "Budget", "rows": [["=1+1"]]}])
    assert len(ALL_TEMP_FILES) == before


def test_create_pptx_allows_max_slides_with_notes():
    from app.services.office_documents import MAX_SLIDES

    slides = [{"title": f"Slide {index}", "bullets": ["Point"], "notes": "Speaker"} for index in range(MAX_SLIDES)]
    result = create_pptx(slides=slides)
    presentation = Presentation(io.BytesIO(base64.b64decode(result.content_base64)))
    assert len(presentation.slides) == MAX_SLIDES
    with pytest.raises(OfficeDocumentError) as caught:
        update_pptx(
            content_base64=result.content_base64,
            append_slides=[{"title": "One more"}],
        )
    assert caught.value.message == f"Too many slides (max {MAX_SLIDES})"


def test_update_docx_replace_once_in_merged_cell():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    anchor = table.cell(0, 0)
    anchor.merge(table.cell(0, 1))
    anchor.text = "a.b"
    buffer = io.BytesIO()
    document.save(buffer)
    original = base64.b64encode(buffer.getvalue()).decode("ascii")
    result = update_docx(
        content_base64=original,
        replacements=[{"find": ".", "replace": ".."}],
    )
    assert result.replacements == 1
    assert "a..b" in _docx_text(result.content_base64)
    assert "a....b" not in _docx_text(result.content_base64)


def test_update_rejects_unsafe_zip_and_too_many_entries():
    unsafe = io.BytesIO()
    with zipfile.ZipFile(unsafe, "w") as package:
        package.writestr("word/document.xml", b"<x/>")
        package.writestr("../evil.xml", b"x")
    with pytest.raises(OfficeDocumentError) as path_error:
        update_docx(
            content_base64=base64.b64encode(unsafe.getvalue()).decode("ascii"),
            append_paragraphs=[{"text": "x"}],
        )
    assert path_error.value.message == "Office package contains an unsafe path"

    from app.services.office_documents import MAX_ZIP_ENTRIES

    crowded = io.BytesIO()
    with zipfile.ZipFile(crowded, "w") as package:
        package.writestr("word/document.xml", b"<x/>")
        for index in range(MAX_ZIP_ENTRIES):
            package.writestr(f"word/extra{index}.xml", b"a")
    with pytest.raises(OfficeDocumentError) as count_error:
        update_docx(
            content_base64=base64.b64encode(crowded.getvalue()).decode("ascii"),
            append_paragraphs=[{"text": "x"}],
        )
    assert count_error.value.message == f"Zip entry count exceeds {MAX_ZIP_ENTRIES}"


def test_slide_and_sheet_validation_messages():
    with pytest.raises(OfficeDocumentError) as slide:
        create_pptx(slides=[{"title": "  ", "bullets": []}])
    assert slide.value.message == "slide needs a title or at least one bullet"
    with pytest.raises(OfficeDocumentError) as sheets:
        create_xlsx(sheets=[])
    assert sheets.value.message == "Provide at least one sheet"
    with pytest.raises(OfficeDocumentError) as value:
        create_xlsx(sheets=[{"name": "Budget", "rows": [[{"nested": 1}]]}])
    assert value.value.message == "value must be a string, number, boolean, or null"
