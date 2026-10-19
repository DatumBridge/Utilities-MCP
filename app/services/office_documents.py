"""Build and edit Word, Excel, and PowerPoint packages in memory.

No disk and no network. Callers upload the returned Base64 with Drive or
Office 365 ``upload_file``. Spreadsheet strings are stored as text so a value
that starts with ``=`` is not an Excel formula.
"""

from __future__ import annotations

import base64
import io
import re
import threading
import zipfile
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.utils.cell import coordinate_to_tuple
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Inches, Pt
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsmap, qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from app.core.exceptions import OfficeDocumentError
from app.services.base64_forms import standard_base64
from app.services.xlsx_package import XlsxPackage

MAX_BYTES = 8 * 1024 * 1024
MAX_BASE64_CHARS = 12_000_000
MAX_UNCOMPRESSED = 32 * 1024 * 1024
# 80 slides with speaker notes is about 360 ZIP parts. Keep headroom above that.
MAX_ZIP_ENTRIES = 512
MAX_TEXT_CHARS = 32_000
MAX_TOTAL_TEXT = 2_000_000
MAX_PARAGRAPHS = 400
MAX_TABLES = 20
MAX_TABLE_ROWS = 1_000
MAX_TABLE_COLS = 40
MAX_SHEETS = 20
MAX_ROWS = 5_000
MAX_SHEET_COLS = 50
MAX_SET_CELLS = 500
MAX_SLIDES = 80
MAX_BULLETS = 30
MAX_REPLACEMENTS = 50

MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_XLSM = "application/vnd.ms-excel.sheet.macroEnabled.12"
MIME_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

_PARTS = {
    "docx": "word/document.xml",
    "xlsx": "xl/workbook.xml",
    "xlsm": "xl/workbook.xml",
    "pptx": "ppt/presentation.xml",
}
_DEFAULT_NAMES = {
    "docx": "document.docx",
    "xlsx": "workbook.xlsx",
    "xlsm": "workbook.xlsm",
    "pptx": "presentation.pptx",
}
_MIMES = {"docx": MIME_DOCX, "xlsx": MIME_XLSX, "xlsm": MIME_XLSM, "pptx": MIME_PPTX}
_STEM = re.compile(r"^[A-Za-z0-9][A-Za-z0-9 ._-]{0,80}$")
_BAD_SHEET = re.compile(r"[\\*?:/\[\]]")
_CELL_REF = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")
_XLSX_SAVE_LOCK = threading.Lock()
_DOCX_NOTE_RELS = (RELATIONSHIP_TYPE.FOOTNOTES, RELATIONSHIP_TYPE.ENDNOTES, RELATIONSHIP_TYPE.COMMENTS)


@dataclass(frozen=True)
class OfficeFile:
    content_base64: str
    file_name: str
    mime_type: str
    byte_length: int
    replacements: int | None = None


def create_docx(
    *,
    file_name: str | None = None,
    title: str | None = None,
    paragraphs: list | None = None,
    tables: list | None = None,
) -> OfficeFile:
    name = _file_name(file_name, "docx")
    budget = _new_budget()
    heading = _optional_text(title, "title")
    blocks = [
        _parse_paragraph(item, budget)
        for item in _items(paragraphs, MAX_PARAGRAPHS, "paragraphs")
    ]
    parsed_tables = [_parse_table(item, budget) for item in _items(tables, MAX_TABLES, "tables")]
    if heading is None and not blocks and not parsed_tables:
        raise OfficeDocumentError(
            "VALIDATION_ERROR",
            "Provide at least one paragraph, table, or title",
        )
    if heading is not None:
        _charge_text(budget, [heading])
    document = Document()
    if heading is not None:
        document.add_heading(heading, level=1)
    for text, level in blocks:
        _add_docx_paragraph(document, text, level)
    for table in parsed_tables:
        _add_docx_table(document, table)
    return _package(document, name, "docx")


def update_docx(
    *,
    content_base64: str,
    file_name: str | None = None,
    append_paragraphs: list | None = None,
    append_tables: list | None = None,
    replacements: list | None = None,
) -> OfficeFile:
    name = _file_name(file_name, "docx")
    budget = _new_budget()
    blocks = [
        _parse_paragraph(item, budget)
        for item in _items(append_paragraphs, MAX_PARAGRAPHS, "paragraphs")
    ]
    parsed_tables = [_parse_table(item, budget) for item in _items(append_tables, MAX_TABLES, "tables")]
    edits = _parse_replacements(replacements)
    if not blocks and not parsed_tables and not edits:
        raise OfficeDocumentError(
            "VALIDATION_ERROR",
            "Provide at least one of append_paragraphs, append_tables, or replacements",
        )
    document = _open_docx(_decode_package(content_base64, "docx"))
    matched = _replace_docx(document, edits) if edits else None
    if edits and matched == 0 and not blocks and not parsed_tables:
        raise OfficeDocumentError("VALIDATION_ERROR", "replacement matched no text")
    for text, level in blocks:
        _add_docx_paragraph(document, text, level)
    for table in parsed_tables:
        _add_docx_table(document, table)
    return _package(document, name, "docx", replacements=matched)


def create_xlsx(*, file_name: str | None = None, sheets: list | None = None) -> OfficeFile:
    name = _file_name(file_name, "xlsx")
    parsed = _parse_sheets(sheets, required=True, budget=_new_budget())
    workbook = Workbook()
    default = workbook.active
    # openpyxl's default sheet is named Sheet. Rename it before create_sheet so a
    # requested title of Sheet or Sheet1 is not silently rewritten.
    default.title = _unused_sheet_title({spec["name"].casefold() for spec in parsed})
    for spec in parsed:
        worksheet = workbook.create_sheet(spec["name"])
        _write_sheet_body(worksheet, spec["headers"], spec["rows"])
    workbook.remove(default)
    return _package(workbook, name, "xlsx")


def update_xlsx(
    *,
    content_base64: str,
    file_name: str | None = None,
    set_cells: list | None = None,
    append_rows: list | None = None,
    add_sheets: list | None = None,
) -> OfficeFile:
    budget = _new_budget()
    cells = _parse_set_cells(set_cells, budget)
    appends = _parse_appends(append_rows, budget)
    sheets = _parse_sheets(add_sheets, required=False, budget=budget)
    if not cells and not appends and not sheets:
        raise OfficeDocumentError(
            "VALIDATION_ERROR",
            "Provide at least one of set_cells, append_rows, or add_sheets",
        )
    package = XlsxPackage(_decode_package(content_base64, "xlsx"))
    kind = "xlsm" if package.macro_enabled else "xlsx"
    name = _file_name(file_name, kind)
    if sheets and package.sheet_count() + len(sheets) > MAX_SHEETS:
        raise OfficeDocumentError("VALIDATION_ERROR", f"Too many sheets (max {MAX_SHEETS})")
    for spec in sheets:
        if package.find(spec["name"]) is not None:
            raise OfficeDocumentError("VALIDATION_ERROR", f"duplicate sheet name: {spec['name']}")
        worksheet = package.add_worksheet(spec["name"])
        rows = ([spec["headers"]] if spec["headers"] else []) + spec["rows"]
        for row_index, row_values in enumerate(rows, start=1):
            for column, value in enumerate(row_values, start=1):
                worksheet.write(row_index, column, value)
    for spec in cells:
        package.worksheet(spec["sheet"]).write(spec["row"], spec["column"], spec["value"])
    for spec in appends:
        worksheet = package.worksheet(spec["sheet"])
        start = worksheet.next_row()
        if start + len(spec["rows"]) - 1 > MAX_ROWS:
            raise OfficeDocumentError("VALIDATION_ERROR", f"Too many rows (max {MAX_ROWS})")
        for offset, row_values in enumerate(spec["rows"]):
            for column, value in enumerate(row_values, start=1):
                worksheet.write(start + offset, column, value)
    return _finish(package.save(), name, kind)


def create_pptx(*, file_name: str | None = None, slides: list | None = None) -> OfficeFile:
    name = _file_name(file_name, "pptx")
    parsed = _parse_slides(slides, required=True, budget=_new_budget())
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    for spec in parsed:
        _add_slide(presentation, spec)
    return _package(presentation, name, "pptx")


def update_pptx(
    *,
    content_base64: str,
    file_name: str | None = None,
    append_slides: list | None = None,
    replacements: list | None = None,
    set_slide_titles: list | None = None,
) -> OfficeFile:
    name = _file_name(file_name, "pptx")
    budget = _new_budget()
    slides = _parse_slides(append_slides, required=False, budget=budget)
    edits = _parse_replacements(replacements)
    titles = _parse_titles(set_slide_titles, budget)
    if not slides and not edits and not titles:
        raise OfficeDocumentError(
            "VALIDATION_ERROR",
            "Provide at least one of append_slides, replacements, or set_slide_titles",
        )
    presentation = _open_pptx(_decode_package(content_base64, "pptx"))
    matched = _replace_pptx(presentation, edits) if edits else None
    if edits and matched == 0 and not slides and not titles:
        raise OfficeDocumentError("VALIDATION_ERROR", "replacement matched no text")
    count = len(presentation.slides)
    for spec in titles:
        if spec["index"] >= count:
            raise OfficeDocumentError("VALIDATION_ERROR", f"slide index out of range: {spec['index']}")
        _set_slide_title(presentation, presentation.slides[spec["index"]], spec["title"])
    if count + len(slides) > MAX_SLIDES:
        raise OfficeDocumentError("VALIDATION_ERROR", f"Too many slides (max {MAX_SLIDES})")
    for spec in slides:
        _add_slide(presentation, spec)
    return _package(presentation, name, "pptx", replacements=matched)


def _file_name(name: str | None, kind: str) -> str:
    extension = f".{kind}"
    if name is None or str(name).strip() == "":
        return _DEFAULT_NAMES[kind]
    if not isinstance(name, str):
        raise OfficeDocumentError(
            "VALIDATION_ERROR",
            f"file_name must be a base name ending with {extension} (letters, numbers, spaces, dots, underscores, hyphens)",
        )
    cleaned = name.strip()
    message = (
        f"file_name must be a base name ending with {extension} "
        "(letters, numbers, spaces, dots, underscores, hyphens)"
    )
    if "/" in cleaned or "\\" in cleaned or ".." in cleaned or not cleaned.lower().endswith(extension):
        raise OfficeDocumentError("VALIDATION_ERROR", message)
    stem = cleaned[: -len(extension)]
    if not _STEM.match(stem):
        raise OfficeDocumentError("VALIDATION_ERROR", message)
    return stem + extension


def _optional_text(value: str | None, label: str) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise OfficeDocumentError("VALIDATION_ERROR", f"{label} must be a string")
    if len(value) > MAX_TEXT_CHARS:
        raise OfficeDocumentError("VALIDATION_ERROR", f"{label} exceeds {MAX_TEXT_CHARS} characters")
    stripped = value.strip()
    return stripped or None


def _require_text(value, label: str) -> str:
    if not isinstance(value, str):
        raise OfficeDocumentError("VALIDATION_ERROR", f"{label} must be a string")
    if len(value) > MAX_TEXT_CHARS:
        raise OfficeDocumentError("VALIDATION_ERROR", f"{label} exceeds {MAX_TEXT_CHARS} characters")
    return value


def _new_budget() -> dict[str, int]:
    return {"n": 0}


def _charge_text(budget: dict[str, int], chunks: list[str]) -> None:
    budget["n"] += sum(len(chunk) for chunk in chunks)
    if budget["n"] > MAX_TOTAL_TEXT:
        raise OfficeDocumentError("VALIDATION_ERROR", f"Document text exceeds {MAX_TOTAL_TEXT} characters")


def _items(values, limit: int, noun: str) -> list:
    if not values:
        return []
    if not isinstance(values, list):
        raise OfficeDocumentError("VALIDATION_ERROR", f"{noun} must be an array")
    if len(values) > limit:
        raise OfficeDocumentError("VALIDATION_ERROR", f"Too many {noun} (max {limit})")
    return values


def _obj(item, label: str) -> dict:
    if hasattr(item, "model_dump"):
        item = item.model_dump()
    if not isinstance(item, dict):
        raise OfficeDocumentError("VALIDATION_ERROR", f"{label} must be an object")
    return item


def _parse_paragraph(item, budget: dict[str, int]) -> tuple[str, int]:
    raw = _obj(item, "paragraph")
    text = _require_text(raw.get("text"), "paragraph text")
    level = raw.get("heading_level", 0)
    if isinstance(level, bool) or level not in (0, 1, 2, 3):
        raise OfficeDocumentError("VALIDATION_ERROR", "heading_level must be 0, 1, 2, or 3")
    _charge_text(budget, [text])
    return text, int(level)


def _parse_table(item, budget: dict[str, int]) -> dict:
    raw = _obj(item, "table")
    headers = raw.get("headers") or []
    rows = raw.get("rows") or []
    if not isinstance(headers, list) or not isinstance(rows, list):
        raise OfficeDocumentError("VALIDATION_ERROR", "table headers and rows must be arrays")
    if not headers and not rows:
        raise OfficeDocumentError("VALIDATION_ERROR", "table needs headers or rows")
    if len(rows) > MAX_TABLE_ROWS:
        raise OfficeDocumentError("VALIDATION_ERROR", f"Too many table rows (max {MAX_TABLE_ROWS})")
    header_cells = [_require_text(cell, "header") for cell in headers]
    parsed_rows = []
    for row in rows:
        if not isinstance(row, list):
            raise OfficeDocumentError("VALIDATION_ERROR", "table row must be an array of strings")
        parsed_rows.append([_require_text(cell, "table cell") for cell in row])
    width = len(header_cells) if header_cells else max(len(row) for row in parsed_rows)
    if width > MAX_TABLE_COLS:
        raise OfficeDocumentError("VALIDATION_ERROR", f"Too many table columns (max {MAX_TABLE_COLS})")
    if width == 0:
        raise OfficeDocumentError("VALIDATION_ERROR", "table needs headers or rows")
    for row in parsed_rows:
        if len(row) > width:
            raise OfficeDocumentError("VALIDATION_ERROR", "table row has more cells than columns")
    _charge_text(budget, header_cells + [cell for row in parsed_rows for cell in row])
    return {"headers": header_cells, "rows": parsed_rows, "width": width}


def _parse_replacements(values) -> list[tuple[str, str]]:
    items = _items(values, MAX_REPLACEMENTS, "replacements")
    edits = []
    for item in items:
        raw = _obj(item, "replacement")
        find = raw.get("find")
        replace = raw.get("replace")
        if not isinstance(find, str) or find == "":
            raise OfficeDocumentError("VALIDATION_ERROR", "find must be a non-empty string")
        if not isinstance(replace, str):
            raise OfficeDocumentError("VALIDATION_ERROR", "replace must be a string")
        if len(find) > MAX_TEXT_CHARS or len(replace) > MAX_TEXT_CHARS:
            raise OfficeDocumentError("VALIDATION_ERROR", f"replacement exceeds {MAX_TEXT_CHARS} characters")
        edits.append((find, replace))
    return edits


def _add_docx_paragraph(document, text: str, level: int) -> None:
    if level == 0:
        document.add_paragraph(text)
        return
    document.add_heading(text, level=level)


def _add_docx_table(document, spec: dict) -> None:
    headers = spec["headers"]
    data_rows = spec["rows"]
    width = spec["width"]
    # Word merges two tables with nothing between them into one table.
    blocks = [child for child in document.element.body if child.tag != qn("w:sectPr")]
    if blocks and blocks[-1].tag == qn("w:tbl"):
        document.add_paragraph()
    table = document.add_table(rows=1, cols=width)
    table.style = "Table Grid"
    if headers:
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        body = data_rows
    else:
        first = data_rows[0]
        for index in range(width):
            table.rows[0].cells[index].text = first[index] if index < len(first) else ""
        body = data_rows[1:]
    for row in body:
        cells = table.add_row().cells
        for index in range(width):
            cells[index].text = row[index] if index < len(row) else ""


def _replace_runs(runs: list, find: str, replace: str) -> int:
    """Replace text across runs. A match goes into the run where it starts; other runs keep their formatting."""
    texts = [run.text or "" for run in runs]
    full = "".join(texts)
    starts = []
    index = full.find(find)
    while index != -1:
        starts.append(index)
        index = full.find(find, index + len(find))
    if not starts:
        return 0
    position = 0
    match = 0
    for run, text in zip(runs, texts):
        pieces = []
        for char in text:
            while match < len(starts) and position >= starts[match] + len(find):
                match += 1
            if match < len(starts) and position >= starts[match]:
                if position == starts[match]:
                    pieces.append(replace)
            else:
                pieces.append(char)
            position += 1
        updated = "".join(pieces)
        if updated != text:
            run.text = updated
    return len(starts)


def _docx_runs(paragraph) -> list:
    # Paragraph.runs skips runs inside hyperlinks, but Paragraph.text includes them.
    return [Run(element, paragraph) for element in paragraph._p.xpath("./w:r | ./w:hyperlink/w:r")]


def _docx_paragraphs(container, seen: set) -> list:
    paragraphs = list(container.paragraphs)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                # Merged cells repeat the anchor cell once per grid slot.
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                paragraphs.extend(_docx_paragraphs(cell, seen))
    return paragraphs


def _docx_header_footers(document) -> list:
    parts = []
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            # A linked part has no definition of its own. Reading it would add one.
            if not part.is_linked_to_previous:
                parts.append(part)
    return parts


def _docx_text_box_paragraphs(root, parent, seen: set) -> list[tuple[object, bool]]:
    """Text box paragraphs. The VML copy under mc:Fallback is edited but not counted."""
    found = []
    for element in root.xpath(".//w:txbxContent//w:p"):
        if element in seen:
            continue
        seen.add(element)
        counted = not element.xpath("ancestor::*[local-name()='Fallback']")
        found.append((Paragraph(element, parent), counted))
    return found


def _docx_note_parts(document) -> list[tuple[object, object]]:
    """Footnotes, endnotes, and comments. python-docx keeps notes as raw bytes, so they are parsed here."""
    found = []
    for rel in document.part.rels.values():
        if rel.is_external or rel.reltype not in _DOCX_NOTE_RELS:
            continue
        part = rel.target_part
        root = part.element if hasattr(part, "element") else parse_xml(part.blob)
        found.append((part, root))
    return found


def _docx_note_paragraphs(root, parent, seen: set) -> list[tuple[object, bool]]:
    found = []
    for element in root.iter(f"{{{nsmap['w']}}}p"):
        if element in seen:
            continue
        seen.add(element)
        counted = not element.xpath("ancestor::*[local-name()='Fallback']")
        found.append((Paragraph(element, parent), counted))
    return found


def _replace_docx(document, edits: list[tuple[str, str]]) -> int:
    seen: set = set()
    parts = _docx_header_footers(document)
    targets = [(paragraph, True, None) for paragraph in _docx_paragraphs(document, seen)]
    for part in parts:
        targets.extend((paragraph, True, None) for paragraph in _docx_paragraphs(part, seen))
    targets.extend((p, counted, None) for p, counted in _docx_text_box_paragraphs(document.element.body, document, seen))
    for part in parts:
        targets.extend((p, counted, None) for p, counted in _docx_text_box_paragraphs(part._element, part, seen))
    for part, root in _docx_note_parts(document):
        targets.extend((p, counted, (part, root)) for p, counted in _docx_note_paragraphs(root, document, seen))
    total = 0
    changed = {}
    for find, replace in edits:
        for paragraph, counted, owner in targets:
            matched = _replace_runs(_docx_runs(paragraph), find, replace)
            if counted:
                total += matched
            if matched and owner is not None:
                changed[id(owner[0])] = owner
    for part, root in changed.values():
        if not hasattr(part, "element"):
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return total


def _sheet_name(value) -> str:
    if not isinstance(value, str):
        raise OfficeDocumentError("VALIDATION_ERROR", "sheet name must be a string")
    cleaned = value.strip()
    if (
        not cleaned
        or len(cleaned) > 31
        or _BAD_SHEET.search(cleaned)
        or cleaned.startswith("'")
        or cleaned.endswith("'")
    ):
        raise OfficeDocumentError(
            "VALIDATION_ERROR",
            "sheet name must be 1-31 characters and must not contain \\ / * ? : [ ] or start or end with a quote",
        )
    return cleaned


def _parse_cell_value(value, budget: dict[str, int]):
    if value is None or isinstance(value, (str, bool, int, float)):
        if isinstance(value, bool):
            return value
        if isinstance(value, float) and (value != value or value in (float("inf"), float("-inf"))):
            raise OfficeDocumentError("VALIDATION_ERROR", "cell value must be a finite number")
        if isinstance(value, str):
            if len(value) > MAX_TEXT_CHARS:
                raise OfficeDocumentError("VALIDATION_ERROR", f"cell text exceeds {MAX_TEXT_CHARS} characters")
            _charge_text(budget, [value])
        return value
    raise OfficeDocumentError("VALIDATION_ERROR", "value must be a string, number, boolean, or null")


def _parse_row(row, budget: dict[str, int]) -> list:
    if not isinstance(row, list):
        raise OfficeDocumentError("VALIDATION_ERROR", "row must be an array")
    if len(row) > MAX_SHEET_COLS:
        raise OfficeDocumentError("VALIDATION_ERROR", f"Too many sheet columns (max {MAX_SHEET_COLS})")
    return [_parse_cell_value(value, budget) for value in row]


def _parse_sheets(values, *, required: bool, budget: dict[str, int]) -> list[dict]:
    if not values:
        if required:
            raise OfficeDocumentError("VALIDATION_ERROR", "Provide at least one sheet")
        return []
    items = _items(values, MAX_SHEETS, "sheets")
    parsed = []
    seen = set()
    for item in items:
        raw = _obj(item, "sheet")
        name = _sheet_name(raw.get("name"))
        key = name.casefold()
        if key in seen:
            raise OfficeDocumentError("VALIDATION_ERROR", f"duplicate sheet name: {name}")
        seen.add(key)
        headers = raw.get("headers") or []
        rows = raw.get("rows") or []
        if not isinstance(headers, list) or not isinstance(rows, list):
            raise OfficeDocumentError("VALIDATION_ERROR", "sheet headers and rows must be arrays")
        if len(rows) > MAX_ROWS:
            raise OfficeDocumentError("VALIDATION_ERROR", f"Too many rows (max {MAX_ROWS})")
        header_cells = [_require_text(cell, "header") for cell in headers]
        _charge_text(budget, header_cells)
        if len(header_cells) > MAX_SHEET_COLS:
            raise OfficeDocumentError("VALIDATION_ERROR", f"Too many sheet columns (max {MAX_SHEET_COLS})")
        parsed_rows = []
        for row in rows:
            parsed_row = _parse_row(row, budget)
            if header_cells and len(parsed_row) > len(header_cells):
                raise OfficeDocumentError("VALIDATION_ERROR", "row has more cells than headers")
            parsed_rows.append(parsed_row)
        parsed.append({"name": name, "headers": header_cells, "rows": parsed_rows})
    return parsed


def _parse_set_cells(values, budget: dict[str, int]) -> list[dict]:
    items = _items(values, MAX_SET_CELLS, "set_cells")
    parsed = []
    for item in items:
        raw = _obj(item, "cell")
        sheet = _sheet_name(raw.get("sheet"))
        cell = raw.get("cell")
        if not isinstance(cell, str) or not _CELL_REF.match(cell.strip().upper()):
            raise OfficeDocumentError("VALIDATION_ERROR", "cell must be an A1 reference such as B2")
        row, column = coordinate_to_tuple(cell.strip().upper())
        if row > MAX_ROWS or column > MAX_SHEET_COLS:
            raise OfficeDocumentError("VALIDATION_ERROR", "cell is outside the sheet limits")
        parsed.append(
            {
                "sheet": sheet,
                "row": row,
                "column": column,
                "value": _parse_cell_value(raw.get("value"), budget),
            }
        )
    return parsed


def _parse_appends(values, budget: dict[str, int]) -> list[dict]:
    items = _items(values, 20, "append_rows")
    parsed = []
    for item in items:
        raw = _obj(item, "append")
        rows = raw.get("rows")
        if not isinstance(rows, list) or not rows:
            raise OfficeDocumentError("VALIDATION_ERROR", "append rows must be a non-empty array")
        if len(rows) > MAX_ROWS:
            raise OfficeDocumentError("VALIDATION_ERROR", f"Too many rows (max {MAX_ROWS})")
        parsed.append({"sheet": _sheet_name(raw.get("sheet")), "rows": [_parse_row(row, budget) for row in rows]})
    return parsed


def _write_cell(cell, value) -> None:
    if isinstance(value, str):
        cell.value = value
        cell.data_type = "s"
        cell.number_format = "@"
        return
    cell.value = value


def _write_row(worksheet, row_index: int, values: list) -> None:
    for column, value in enumerate(values, start=1):
        _write_cell(worksheet.cell(row=row_index, column=column), value)


def _write_sheet_body(worksheet, headers: list[str], rows: list[list]) -> None:
    next_row = 1
    if headers:
        _write_row(worksheet, 1, headers)
        next_row = 2
    for offset, row in enumerate(rows):
        _write_row(worksheet, next_row + offset, row)


def _parse_slides(values, *, required: bool, budget: dict[str, int]) -> list[dict]:
    if not values:
        if required:
            raise OfficeDocumentError("VALIDATION_ERROR", "Provide at least one slide")
        return []
    items = _items(values, MAX_SLIDES, "slides")
    parsed = []
    for item in items:
        raw = _obj(item, "slide")
        title = raw.get("title") or ""
        if not isinstance(title, str):
            raise OfficeDocumentError("VALIDATION_ERROR", "title must be a string")
        title = title.strip()
        if len(title) > MAX_TEXT_CHARS:
            raise OfficeDocumentError("VALIDATION_ERROR", f"title exceeds {MAX_TEXT_CHARS} characters")
        bullets = raw.get("bullets") or []
        if not isinstance(bullets, list):
            raise OfficeDocumentError("VALIDATION_ERROR", "bullets must be an array")
        if len(bullets) > MAX_BULLETS:
            raise OfficeDocumentError("VALIDATION_ERROR", f"Too many bullets (max {MAX_BULLETS})")
        bullet_text = [_require_text(bullet, "bullet") for bullet in bullets]
        notes = raw.get("notes")
        if notes is not None:
            notes = _require_text(notes, "notes")
        if not title and not bullet_text:
            raise OfficeDocumentError("VALIDATION_ERROR", "slide needs a title or at least one bullet")
        _charge_text(budget, [title, notes or "", *bullet_text])
        parsed.append({"title": title, "bullets": bullet_text, "notes": notes})
    return parsed


def _parse_titles(values, budget: dict[str, int]) -> list[dict]:
    items = _items(values, MAX_SLIDES, "set_slide_titles")
    parsed = []
    for item in items:
        raw = _obj(item, "title")
        index = raw.get("index")
        if isinstance(index, bool) or not isinstance(index, int) or index < 0:
            raise OfficeDocumentError("VALIDATION_ERROR", "slide index must be a non-negative integer")
        title = raw.get("title")
        if not isinstance(title, str) or title.strip() == "":
            raise OfficeDocumentError("VALIDATION_ERROR", "title must be a non-empty string")
        if len(title) > MAX_TEXT_CHARS:
            raise OfficeDocumentError("VALIDATION_ERROR", f"title exceeds {MAX_TEXT_CHARS} characters")
        cleaned = title.strip()
        _charge_text(budget, [cleaned])
        parsed.append({"index": index, "title": cleaned})
    return parsed


def _blank_layout(presentation):
    for layout in presentation.slide_layouts:
        if (layout.name or "").strip().lower() in ("blank", "blank slide"):
            return layout
    return min(presentation.slide_layouts, key=lambda layout: len(list(layout.placeholders)))


def _add_slide(presentation, spec: dict) -> None:
    slide = presentation.slides.add_slide(_blank_layout(presentation))
    width = presentation.slide_width or Inches(13.333333)
    left = Inches(0.5)
    box_width = max(int(width - Inches(1)), int(Inches(4)))
    title_box = slide.shapes.add_textbox(left, Inches(0.35), Emu(box_width), Inches(1.15))
    title_box.name = "Title"
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.text = spec["title"]
    if title_paragraph.runs:
        title_paragraph.runs[0].font.size = Pt(32)
        title_paragraph.runs[0].font.bold = True
    body = slide.shapes.add_textbox(left, Inches(1.7), Emu(box_width), Inches(5.2))
    body.name = "Body"
    body_frame = body.text_frame
    body_frame.word_wrap = True
    bullets = spec["bullets"]
    if bullets:
        body_frame.paragraphs[0].text = bullets[0]
        if body_frame.paragraphs[0].runs:
            body_frame.paragraphs[0].runs[0].font.size = Pt(20)
        for bullet in bullets[1:]:
            paragraph = body_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            if paragraph.runs:
                paragraph.runs[0].font.size = Pt(20)
    if spec["notes"]:
        slide.notes_slide.notes_text_frame.text = spec["notes"]


def _title_shape(slide):
    for shape in _walk_shapes(slide.shapes):
        if getattr(shape, "name", "") == "Title" and getattr(shape, "has_text_frame", False):
            return shape
    placeholder = slide.shapes.title
    if placeholder is not None and placeholder.has_text_frame:
        return placeholder
    return None


def _set_slide_title(presentation, slide, title: str) -> None:
    target = _title_shape(slide)
    if target is None:
        width = presentation.slide_width or Inches(13.333333)
        target = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Emu(max(int(width - Inches(1)), int(Inches(4)))), Inches(1.15))
        target.name = "Title"
    frame = target.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = title
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = title
    for extra in frame.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def _walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)


def _replace_text_frame(frame, find: str, replace: str) -> int:
    total = 0
    for paragraph in frame.paragraphs:
        runs = list(paragraph.runs)
        if runs:
            total += _replace_runs(runs, find, replace)
            continue
        full = paragraph.text or ""
        if find in full:
            paragraph.text = full.replace(find, replace)
            total += full.count(find)
    return total


def _replace_pptx(presentation, edits: list[tuple[str, str]]) -> int:
    total = 0
    frames = []
    for slide in presentation.slides:
        for shape in _walk_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                frames.append(shape.text_frame)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        frames.append(cell.text_frame)
        if slide.has_notes_slide:
            frames.append(slide.notes_slide.notes_text_frame)
    for find, replace in edits:
        for frame in frames:
            total += _replace_text_frame(frame, find, replace)
    return total


def _decode_package(content_base64: str, kind: str) -> bytes:
    if not isinstance(content_base64, str) or content_base64.strip() == "":
        raise OfficeDocumentError("VALIDATION_ERROR", "content_base64 is required")
    compact = standard_base64("".join(content_base64.split()))
    if len(compact) > MAX_BASE64_CHARS:
        raise OfficeDocumentError("VALIDATION_ERROR", f"File exceeds {MAX_BYTES} bytes")
    try:
        raw = base64.b64decode(compact, validate=True)
    except Exception as exc:
        raise OfficeDocumentError("INVALID_BASE64", f"Invalid Base64: {exc}") from exc
    if len(raw) > MAX_BYTES:
        raise OfficeDocumentError("VALIDATION_ERROR", f"File exceeds {MAX_BYTES} bytes")
    _assert_ooxml(raw, kind, cap_entries=True)
    return raw


def _assert_ooxml(raw: bytes, kind: str, *, cap_entries: bool) -> None:
    if len(raw) < 4 or raw[:2] != b"PK":
        raise OfficeDocumentError("INVALID_OFFICE_FILE", "File is not an Office Open XML package")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as package:
            infos = package.infolist()
            if cap_entries and len(infos) > MAX_ZIP_ENTRIES:
                raise OfficeDocumentError("VALIDATION_ERROR", f"Zip entry count exceeds {MAX_ZIP_ENTRIES}")
            total = 0
            for info in infos:
                filename = info.filename
                if filename.startswith("/") or ".." in filename.replace("\\", "/").split("/"):
                    raise OfficeDocumentError("INVALID_OFFICE_FILE", "Office package contains an unsafe path")
                if info.file_size < 0 or info.file_size > MAX_UNCOMPRESSED:
                    raise OfficeDocumentError("VALIDATION_ERROR", "Uncompressed Office package is too large")
                total += info.file_size
                if total > MAX_UNCOMPRESSED:
                    raise OfficeDocumentError("VALIDATION_ERROR", "Uncompressed Office package is too large")
            names = {info.filename for info in infos}
    except OfficeDocumentError:
        raise
    except zipfile.BadZipFile as exc:
        raise OfficeDocumentError("INVALID_OFFICE_FILE", "Corrupt Office package") from exc
    part = _PARTS[kind]
    if part not in names:
        raise OfficeDocumentError("INVALID_OFFICE_FILE", f"Package is missing {part}")


def _open_docx(raw: bytes):
    try:
        return Document(io.BytesIO(raw))
    except OfficeDocumentError:
        raise
    except Exception as exc:
        raise OfficeDocumentError("INVALID_OFFICE_FILE", "Could not read Word document") from exc


def _open_pptx(raw: bytes):
    try:
        return Presentation(io.BytesIO(raw))
    except OfficeDocumentError:
        raise
    except Exception as exc:
        raise OfficeDocumentError("INVALID_OFFICE_FILE", "Could not read PowerPoint presentation") from exc


def _unused_sheet_title(reserved: set[str]) -> str:
    candidate = "_tmp"
    number = 0
    while candidate.casefold() in reserved:
        number += 1
        candidate = f"_tmp{number}"
        if number > MAX_SHEETS + 1:
            raise OfficeDocumentError("VALIDATION_ERROR", f"Too many sheets (max {MAX_SHEETS})")
    return candidate


def _save_workbook_bytes(workbook, buffer: io.BytesIO) -> None:
    """Save without openpyxl's NamedTemporaryFile worksheet writer."""
    from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
    from openpyxl.worksheet._writer import WorksheetWriter
    from openpyxl.writer.excel import ExcelWriter, save_workbook

    original = ExcelWriter.write_worksheet

    def write_worksheet(self, ws):
        if self.workbook.write_only:
            raise OfficeDocumentError("INVALID_OFFICE_FILE", "Could not build Excel workbook")
        ws._drawing = SpreadsheetDrawing()
        ws._drawing.charts = ws._charts
        ws._drawing.images = ws._images
        sheet_xml = io.BytesIO()
        writer = WorksheetWriter(ws, out=sheet_xml)
        writer.write()
        self._archive.writestr(ws.path[1:], sheet_xml.getvalue())
        self.manifest.append(ws)
        ws._rels = writer._rels

    with _XLSX_SAVE_LOCK:
        ExcelWriter.write_worksheet = write_worksheet
        try:
            save_workbook(workbook, buffer)
        finally:
            ExcelWriter.write_worksheet = original


def _package(document, file_name: str, kind: str, replacements: int | None = None) -> OfficeFile:
    buffer = io.BytesIO()
    try:
        if kind == "xlsx":
            _save_workbook_bytes(document, buffer)
        else:
            document.save(buffer)
    except OfficeDocumentError:
        raise
    except Exception as exc:
        label = {"docx": "Word document", "xlsx": "Excel workbook", "pptx": "PowerPoint presentation"}[kind]
        raise OfficeDocumentError("INVALID_OFFICE_FILE", f"Could not build {label}") from exc
    return _finish(buffer.getvalue(), file_name, kind, replacements)


def _finish(raw: bytes, file_name: str, kind: str, replacements: int | None = None) -> OfficeFile:
    if len(raw) > MAX_BYTES:
        raise OfficeDocumentError("VALIDATION_ERROR", f"File exceeds {MAX_BYTES} bytes")
    _assert_ooxml(raw, kind, cap_entries=False)
    return OfficeFile(
        content_base64=base64.b64encode(raw).decode("ascii"),
        file_name=file_name,
        mime_type=_MIMES[kind],
        byte_length=len(raw),
        replacements=replacements,
    )
